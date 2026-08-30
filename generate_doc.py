import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = docx.Document()

for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def add_title(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = RGBColor(18, 18, 18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p

def add_h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(237, 30, 58)
    return p

def add_h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 30, 30)
    return p

def add_body(text, bold_prefix=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(20, 20, 20)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(60, 60, 60)
    return p

def add_code(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.2)
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(40, 40, 40)
    return p

add_title("Manxel Full-Stack Developer Assessment\nProject Explanation Document")

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_sub = p_sub.add_run("Candidate Technical Documentation • August 2026")
r_sub.font.name = 'Calibri'
r_sub.font.size = Pt(12)
r_sub.font.color.rgb = RGBColor(120, 120, 120)

doc.add_paragraph()

add_h1("1. Project Overview")
add_body("This project was developed as a complete full-stack web application for Manxel Studio according to the requirements specified in the technical assessment. The solution comprises a modern, high-performance React.js frontend interface paired with a secure Node.js and Express.js REST API backend, backed by MongoDB for persistent data management.")
add_body("The primary goal was to create an engaging digital presence for Manxel showcasing its core brand competencies—Visual Identity, Web Design, and App Design—accompanied by a fully functional client contact pipeline with automated email notifications.")

add_h2("Technologies Used")
add_body("React 19, Vite 8, Tailwind CSS v4, Framer Motion, Axios, React Icons", "• Frontend: ")
add_body("Node.js, Express.js, CORS, Dotenv, Nodemailer", "• Backend: ")
add_body("MongoDB Community Server / Atlas, Mongoose ODM", "• Database: ")

add_h2("Technology Selection Rationale")
add_body("React's component-based paradigm enables high code reusability, clean separation of concerns, and declarative UI rendering. Vite was selected as the build tool to provide instantaneous Hot Module Replacement (HMR) and optimized rollup production bundles.", "• Why React.js? ")
add_body("Express provides a robust, lightweight, and unopinionated REST API framework with asynchronous request processing, making it ideal for handling client inquiries without thread blocking.", "• Why Node.js / Express? ")
add_body("Inquiry submissions are naturally structured as self-contained documents. MongoDB offers a flexible JSON-like schema with Mongoose validation, eliminating database schema migration bottlenecks and ensuring high write throughput.", "• Why MongoDB? ")

add_h1("2. System Architecture & Complete Data Flow")
add_body("The end-to-end communication follows standard RESTful principles with dual-tier validation:")
add_code("React Form (Contact.jsx)\n  └──> [Frontend Validation: Regex & Required Checks]\n       └──> [Axios HTTP POST /api/contact (Vite Proxy:5000)]\n            └──> [Express Route: routes/contact.js]\n                 └──> [Controller: contactController.js]\n                      ├──> [Server Validation & Sanitization]\n                      ├──> [MongoDB Model: Contact.create()]\n                      ├──> [Nodemailer: Async Gmail Lead Dispatch]\n                      └──> [HTTP 201 JSON Response to Client UI]")

add_h1("3. Frontend Explanation & Component Hierarchy")
add_body("The frontend application is structured into modular, single-responsibility components under src/components/:")
add_body("Fixed header featuring responsive mobile navigation drawer, brand identity badge, and smooth scroll jump links.", "• Navbar.jsx: ")
add_body("Dynamic hero banner with high-impact typography, action CTAs, and a continuous moving preview showcase.", "• Hero.jsx: ")
add_body("Strategic company breakdown featuring 6 Core Value cards with visual dot rating progression and an interactive Process section with dynamic dark mode transition on hover.", "• About.jsx: ")
add_body("What We Do section presenting 3 two-tier service cards (Visual Identity, Web Design, App Design) with dual-tone SVGs, tag badges, and imagery.", "• Services.jsx: ")
add_body("8 curated showcase project cards featuring hover zoom, blur depth, and slide-up link action bars.", "• Portfolio.jsx: ")
add_body("Continuous 6-logo SVG client marquee with seamless zero-glitch infinite loop, accompanied by an interactive 5-item FAQ accordion.", "• FAQ.jsx: ")
add_body("Two-column contact layout featuring direct contact cards and an interactive dark form with real-time field validation, loading states, and error handling.", "• Contact.jsx: ")
add_body("Solid #0d0d0d background footer with atmospheric red glow, navigation links, and a subtle MANXEL watermark.", "• Footer.jsx: ")

add_h1("4. Backend Implementation & API Details")
add_body("The backend server operates on port 5000 and exposes REST endpoints:")
add_body("POST /api/contact — Receives form submission payload (name, email, phone, subject, message), executes validation, inserts record into MongoDB, and dispatches an HTML email alert via Nodemailer.", "• Contact Endpoint: ")
add_body("GET /api/health — Returns server health status and ISO timestamp for monitoring.", "• Health Endpoint: ")

add_h2("Database Model (Mongoose Schema)")
add_code("""const contactSchema = new mongoose.Schema({
  name: { type: String, required: true, trim: true, maxlength: 100 },
  email: { type: String, required: true, trim: true, match: /^[^\s@]+@[^\s@]+\.[^\s@]+$/ },
  phone: { type: String, required: true, trim: true, maxlength: 20 },
  subject: { type: String, required: true, trim: true, maxlength: 200 },
  message: { type: String, required: true, trim: true, maxlength: 5000 }
}, { timestamps: true });""")

add_h1("5. Setup & Running Instructions")
add_body("1. Clone or extract the project repository.")
add_body("2. Copy server/.env.example to server/.env and configure PORT and MONGO_URI.")
add_body("3. Install dependencies: run 'npm install' in root, and 'cd server && npm install'.")
add_body("4. Start backend: run 'cd server && npm start'. (Server runs on http://localhost:5000)")
add_body("5. Start frontend: in a new terminal, run 'npm run dev'. (App runs on http://localhost:5173)")

add_h1("6. Submission Deliverables Checklist")
add_body("✓ Complete React.js Frontend source code and build assets.")
add_body("✓ Complete Node.js / Express.js Backend source code.")
add_body("✓ MongoDB Schema & Mongoose Model integration.")
add_body("✓ Environment configuration template (.env.example).")
add_body("✓ Project README.md documentation.")
add_body("✓ Project Explanation Document (Project-Documentation.docx).")

doc.save("c:/Users/kk312/OneDrive/Desktop/Manxel/Project-Documentation.docx")
print("SUCCESS: Project-Documentation.docx generated successfully!")
